import io
import json
import os
import sys
import time

import requests
from openpyxl import load_workbook

BASE = os.environ["UAT_API_BASE_URL"].rstrip("/")
PASSWORD = os.environ["UAT_DEMO_PASSWORD"]
DATE = "2026-10-01"

CENTER_CODE = "UAT-C-TEAM003"
SITE_CODE = "UAT-S-TEAM003"
TEAM_CODE = "TEAM-003"
EMPLOYEE_NO = "pilotv01"
LOCATION_CODE = "LOC-260812-696991"


def request(method, path, **kwargs):
    last = None
    for attempt in range(1, 6):
        try:
            r = requests.request(method, BASE + path, timeout=60, **kwargs)
            r.raise_for_status()
            return r
        except requests.RequestException as ex:
            last = ex
            if attempt == 5:
                raise
            print(f"TRANSIENT_HTTP_RETRY path={path} attempt={attempt}/5")
            time.sleep(5)
    raise last


def login(account):
    return request(
        "POST",
        "/api/v1/auth/demo-login",
        json={"account": account, "password": PASSWORD},
    ).json()


admin_login = login("pilota01")
admin_token = admin_login["accessToken"]
admin_headers = {
    "Authorization": f"Bearer {admin_token}",
    "X-Active-Role": "admin",
}


def get_json(path):
    return request("GET", path, headers=admin_headers).json()


targets = {
    "employment-status": lambda x: x.get("key") == EMPLOYEE_NO,
    "centers": lambda x: x.get("key") == CENTER_CODE,
    "team-centers": lambda x: x.get("key") == TEAM_CODE and x.get("parentKey") == CENTER_CODE,
    "deployment-sites": lambda x: x.get("key") == SITE_CODE,
    "team-sites": lambda x: x.get("key") == TEAM_CODE and x.get("parentKey") == SITE_CODE,
    "employment-sites": lambda x: x.get("key") == EMPLOYEE_NO and x.get("parentKey") == SITE_CODE,
}

for kind, match in targets.items():
    rows = get_json(f"/api/v1/admin/master-data/{kind}")
    if any(match(x) for x in rows):
        raise RuntimeError("TARGET_ALREADY_EXISTS:" + kind)

teams = get_json("/api/v1/teams")
team = next((x for x in teams if x.get("teamCode") == TEAM_CODE and x.get("teamId") == 7), None)
if not team:
    raise RuntimeError("TEAM_003_NOT_FOUND")

people = get_json(
    "/api/v1/admin/people?keyword=pilotv01&userType=Internal&role=visitor&page=1&pageSize=20&sort=code_asc"
).get("items", [])
visitor = next((x for x in people if x.get("employeeNo") == EMPLOYEE_NO), None)
if not visitor or visitor.get("primaryTeamId") != 7:
    raise RuntimeError("PILOTV01_TEAM_MISMATCH")

locations = get_json(
    "/api/v1/managed-locations/search?q=%E5%93%A1%E6%9E%97&teamId=7&isActive=true&page=1&pageSize=20"
).get("items", [])
location = next(
    (
        x
        for x in locations
        if x.get("locationCode") == LOCATION_CODE
        and x.get("locationId") == 25
        and x.get("teamId") == 7
        and x.get("isActive") is True
        and x.get("approvalStatus") == "Approved"
    ),
    None,
)
if not location:
    raise RuntimeError("UAT_LOCATION_NOT_ELIGIBLE")

print("PREFLIGHT_TEAM=" + json.dumps(team, ensure_ascii=False, sort_keys=True))
print(
    "PREFLIGHT_VISITOR="
    + json.dumps(
        {
            "userId": visitor.get("userId"),
            "employeeNo": visitor.get("employeeNo"),
            "primaryTeamId": visitor.get("primaryTeamId"),
            "primaryTeamName": visitor.get("primaryTeamName"),
        },
        ensure_ascii=False,
        sort_keys=True,
    )
)
print(
    "PREFLIGHT_LOCATION="
    + json.dumps(
        {
            "locationId": location.get("locationId"),
            "locationCode": location.get("locationCode"),
            "locationName": location.get("locationName"),
            "teamId": location.get("teamId"),
        },
        ensure_ascii=False,
        sort_keys=True,
    )
)

template = request(
    "GET",
    "/api/v1/admin/master-data/bulk/template.xlsx",
    headers=admin_headers,
).content

wb = load_workbook(io.BytesIO(template))
expected = [
    "EmploymentStatus",
    "Centers",
    "TeamCenters",
    "DeploymentSites",
    "TeamSites",
    "EmploymentSites",
]
if wb.sheetnames != expected:
    raise RuntimeError(f"UNEXPECTED_TEMPLATE_SHEETS:{wb.sheetnames}")

headers = {
    "EmploymentStatus": ["EmployeeNo", "Status", "EffectiveFrom", "EffectiveTo"],
    "Centers": ["CenterCode", "CenterName", "EffectiveFrom", "EffectiveTo", "IsActive"],
    "TeamCenters": ["TeamCode", "CenterCode", "EffectiveFrom", "EffectiveTo"],
    "DeploymentSites": [
        "CenterCode",
        "SiteCode",
        "SiteName",
        "LocationCode",
        "EffectiveFrom",
        "EffectiveTo",
        "IsActive",
    ],
    "TeamSites": ["TeamCode", "SiteCode", "EffectiveFrom", "EffectiveTo"],
    "EmploymentSites": ["EmployeeNo", "SiteCode", "IsPrimary", "EffectiveFrom", "EffectiveTo"],
}
for sheet, expected_headers in headers.items():
    actual = [c.value for c in wb[sheet][1]]
    if actual != expected_headers:
        raise RuntimeError(f"UNEXPECTED_HEADERS:{sheet}:{actual}")
    ws = wb[sheet]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

wb["EmploymentStatus"].append([EMPLOYEE_NO, "Active", DATE, ""])
wb["Centers"].append([CENTER_CODE, "UAT 員林中心", DATE, "", "true"])
wb["TeamCenters"].append([TEAM_CODE, CENTER_CODE, DATE, ""])
wb["DeploymentSites"].append(
    [CENTER_CODE, SITE_CODE, "UAT 員林就業中心", LOCATION_CODE, DATE, "", "true"]
)
wb["TeamSites"].append([TEAM_CODE, SITE_CODE, DATE, ""])
wb["EmploymentSites"].append([EMPLOYEE_NO, SITE_CODE, "true", DATE, ""])

buffer = io.BytesIO()
wb.save(buffer)
workbook = buffer.getvalue()
print("WORKBOOK_ROWS=6")

preview = request(
    "POST",
    "/api/v1/admin/master-data/bulk/preview",
    headers=admin_headers,
    files={
        "file": (
            "p2b-pilotv01.xlsx",
            workbook,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    },
).json()
print("E2_PREVIEW=" + json.dumps(preview, ensure_ascii=False, sort_keys=True))

if (
    preview.get("totalCount") != 6
    or preview.get("validCount") != 6
    or preview.get("errorCount") != 0
):
    raise RuntimeError("E2_PREVIEW_NOT_6_VALID")
items = preview.get("items") or []
if len(items) != 6 or any(
    x.get("status") != "Valid" or x.get("action") != "Create" for x in items
):
    raise RuntimeError("E2_PREVIEW_ACTION_NOT_ALL_CREATE")

batch_id = preview["importBatchId"]
confirm = request(
    "POST",
    f"/api/v1/admin/master-data/bulk/{batch_id}/confirm",
    headers=admin_headers,
).json()
print("E2_CONFIRM=" + json.dumps(confirm, ensure_ascii=False, sort_keys=True))

if (
    confirm.get("created") != 6
    or confirm.get("updated") != 0
    or confirm.get("unchanged") != 0
    or confirm.get("failed") != 0
    or (confirm.get("errors") or [])
):
    raise RuntimeError("E2_CONFIRM_UNEXPECTED_RESULT")

readiness = get_json("/api/v1/admin/master-data/readiness")
print("MASTER_DATA_READINESS_AFTER=" + json.dumps(readiness, ensure_ascii=False, sort_keys=True))

visitor_login = login("pilotv01")
visitor_headers = {
    "Authorization": f"Bearer {visitor_login['accessToken']}",
    "X-Active-Role": "visitor",
}
context = request(
    "GET",
    "/api/v1/trips/context?visitDate=2026-10-01&teamId=7",
    headers=visitor_headers,
).json()
print("PILOTV01_TRIP_CONTEXT_AFTER=" + json.dumps(context, ensure_ascii=False, sort_keys=True))

if context.get("eligibleForTrip") is not True:
    raise RuntimeError("PILOTV01_CONTEXT_NOT_ELIGIBLE:" + str(context.get("validationCode")))
if context.get("selectedTeamId") != 7:
    raise RuntimeError("PILOTV01_CONTEXT_TEAM_MISMATCH")
if not context.get("eligibleDeploymentSites"):
    raise RuntimeError("PILOTV01_NO_ELIGIBLE_DEPLOYMENT_SITE")
if context.get("defaultStartDeploymentSiteId") is None or context.get("defaultEndDeploymentSiteId") is None:
    raise RuntimeError("PILOTV01_DEFAULT_SITE_MISSING")

print("PILOTV01_MASTER_DATA_ENABLEMENT=SUCCESS")
